from docx import Document

doc=Document()

doc.add_paragraph("DocuLens AI Test Document")
doc.add_paragraph("Employee Name: Rahul Kumar")
doc.add_paragraph("The annual salary is ₹8,40,000.")
doc.add_paragraph("The probation period is 6 months.")
doc.add_paragraph("This document is used to test document extraction.")

doc.save("../sample_documents/test_document.docx")

print("DOCX created successfully")